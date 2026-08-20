#!/usr/bin/env python3
"""Build the RevTrace FGCS submission staging package.

The package is deliberately marked STAGING until the senior author returns
the three corrected figures and the corresponding-author phone is supplied.
"""

from __future__ import annotations

import argparse
from datetime import date
import hashlib
import json
from pathlib import Path
import shutil
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "Paper"
OUT = PAPER / "submission" / "FGCS_RevTrace_STAGING"


TITLE = (
    "RevTrace: Dual-Granularity Feedback Transactions for "
    "State-Certified Streaming ASR Serving"
)
ACK = (
    "The computations in this work were supported by the High Performance "
    "Computing Platform of BUCT (Beijing University of Chemical Technology). "
    "This work received no specific grant from any funding agency in the "
    "public, commercial, or not-for-profit sectors."
)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def copy_file(src: Path, dst: Path) -> None:
    if not src.is_file():
        raise FileNotFoundError(src)
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dst)


def set_font(run, *, size: float, bold: bool = False, color: str = "000000") -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1


def build_cover_letter(path: Path) -> None:
    doc = Document()
    style_doc(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("COVER LETTER")
    set_font(r, size=21, bold=True, color="173F8A")
    p.paragraph_format.space_after = Pt(4)

    p = doc.add_paragraph()
    r = p.add_run(TITLE)
    set_font(r, size=13.5, bold=True)
    p.paragraph_format.space_after = Pt(14)

    for text in (
        date.today().strftime("%d %B %Y"),
        "Professor Michela Taufer",
        "Editor-in-Chief",
        "Future Generation Computer Systems",
    ):
        p = doc.add_paragraph(text)
        p.paragraph_format.space_after = Pt(1)

    doc.add_paragraph()
    doc.add_paragraph("Dear Professor Taufer,")
    doc.add_paragraph(
        "Please consider our manuscript, entitled \"" + TITLE +
        "\", for publication as a regular research article in Future "
        "Generation Computer Systems."
    )
    doc.add_paragraph(
        "Revision-generative ASR services feed generated text into later "
        "rounds, turning a displayed output into a write to future model "
        "state. We formalize the resulting feedback-state contamination "
        "channel and introduce Dual-Granularity Feedback Transactions (DGFT), "
        "which separate low-latency display permission from certificate-bound "
        "state-write permission."
    )
    doc.add_paragraph(
        "The work contributes a state-transaction abstraction, proofs of "
        "uncertified-ancestry exclusion and fork-free committed histories, and "
        "a serving runtime evaluated across Qwen3-ASR and Whisper families. "
        "The experiments include two RTX 3090 GPUs and an independent "
        "single-node eight-V100 contract. RevTrace reaches up to 6.88x "
        "wall-clock speedup while maintaining the declared family-specific "
        "quality contracts and records no observed transaction-install "
        "violation in the reported traces."
    )
    doc.add_paragraph(
        "We believe the manuscript fits FGCS because it combines formal "
        "state-consistency guarantees with high-performance serving, "
        "single-node accelerator scale-out, reproducible system evidence, and "
        "explicit capacity and trust boundaries."
    )
    doc.add_paragraph(
        "This manuscript is original, has not been published previously, and "
        "is not under consideration by another journal. All authors have "
        "approved the manuscript and agree with its submission. The authors "
        "declare no competing interests. Source code and frozen audited result "
        "artifacts are available at https://github.com/zhouz4884-coder/RevTrace."
    )
    doc.add_paragraph("Thank you for your consideration.")
    doc.add_paragraph("Sincerely,")
    p = doc.add_paragraph()
    for text, bold in (
        ("Xi Cao\n", True),
        ("Corresponding author\n", False),
        ("Beijing University of Chemical Technology\n", False),
        ("Email: caoxi@buct.edu.cn", False),
    ):
        r = p.add_run(text)
        set_font(r, size=11, bold=bold)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def build_title_page(src: Path, dst: Path, phone: str | None) -> None:
    doc = Document(src)
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("This work received no specific grant"):
            paragraph.text = ACK
        if "Postal address:" in paragraph.text:
            phone_line = (
                "Corresponding-author phone: " +
                (phone if phone else "[REQUIRED BEFORE SUBMISSION]")
            )
            paragraph.text = paragraph.text.replace(
                "\nPostal address:", f"\n{phone_line}\nPostal address:"
            )
    dst.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dst)


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.rstrip() + "\n", encoding="utf-8")


def build_source_tree() -> None:
    source = OUT / "06_latex_source"
    source.mkdir(parents=True, exist_ok=True)
    for name in ("main.tex", "refs.bib", "elsarticle-num.bst", "lineno.sty"):
        copy_file(PAPER / name, source / name)
    for dirname in ("sections", "tables"):
        target = source / dirname
        target.mkdir(parents=True, exist_ok=True)
        for path in sorted((PAPER / dirname).glob("*.tex")):
            copy_file(path, target / path.name)
    figures = [
        "f1_motivation_senior.png",
        "f2_architecture_senior.jpg",
        "f3_certification_senior.jpg",
        "f3_authority_frontier_column.pdf",
        "f4_production.pdf",
        "f5_quality_column.pdf",
        "f7_v100_dual_plane.pdf",
    ]
    for name in figures:
        copy_file(PAPER / "figures" / name, source / "figures" / name)
    zip_path = OUT / "06_latex_source.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(source.rglob("*")):
            if path.is_file():
                archive.write(path, path.relative_to(source.parent))


def build_package(phone: str | None) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    copy_file(PAPER / "main.pdf", OUT / "01_manuscript" / "RevTrace_FGCS_manuscript.pdf")
    copy_file(
        PAPER / "submission" / "title_page_revtrace.docx",
        OUT / "02_title_page" / "title_page_source_original.docx",
    )
    build_title_page(
        PAPER / "submission" / "title_page_revtrace.docx",
        OUT / "02_title_page" / "RevTrace_title_page.docx",
        phone,
    )
    build_cover_letter(OUT / "03_cover_letter" / "RevTrace_cover_letter.docx")

    highlights = [
        line.strip()
        for line in (PAPER / "submission" / "highlights.txt")
        .read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    if not 3 <= len(highlights) <= 5:
        raise ValueError("FGCS highlights must contain 3--5 bullets")
    if any(len(line) > 85 for line in highlights):
        raise ValueError("FGCS highlight exceeds 85 characters")
    write_text(
        OUT / "04_highlights" / "RevTrace_highlights.txt",
        "\n".join(f"- {line}" for line in highlights),
    )

    for number, source_name in (
        (1, "f1_motivation_senior.png"),
        (2, "f2_architecture_senior.jpg"),
        (3, "f3_certification_senior.jpg"),
        (4, "f3_authority_frontier_column.pdf"),
        (5, "f4_production.pdf"),
        (6, "f5_quality_column.pdf"),
        (7, "f7_v100_dual_plane.pdf"),
    ):
        ext = Path(source_name).suffix.lower()
        copy_file(
            PAPER / "figures" / source_name,
            OUT / "05_figures" / f"Figure_{number}{ext}",
        )
    copy_file(
        PAPER / "figures" / "SENIOR_FIGURE_CHANGE_REQUEST.md",
        OUT / "05_figures" / "FIGURE_CORRECTIONS_PENDING.md",
    )
    copy_file(
        PAPER / "figures" / "SENIOR_FIGURE_QA.md",
        OUT / "05_figures" / "FIGURE_QA.md",
    )

    build_source_tree()
    audit_dir = OUT / "07_audits"
    for name in ("REPORT.md", "cjk_coverage_replay.json", "denominator_audit.json", "contract_inventory.json"):
        copy_file(PAPER / "audits" / name, audit_dir / name)
    copy_file(
        PAPER / "submission" / "PRESUBMISSION_QA.md",
        audit_dir / "PRESUBMISSION_QA.md",
    )

    write_text(
        OUT / "08_declarations" / "declarations.txt",
        """Competing interests
The authors declare that they have no known competing financial interests or personal relationships that could have appeared to influence the work reported in this paper.

Generative AI declaration
OpenAI ChatGPT was used as an auxiliary tool for language editing and LaTeX formatting assistance. It was not used to generate experimental measurements. All technical ideas, implementations, numerical results, claims, figures, and final manuscript decisions were independently reviewed and verified by the authors.

Acknowledgments
The computations in this work were supported by the High Performance Computing Platform of BUCT (Beijing University of Chemical Technology). This work received no specific grant from any funding agency in the public, commercial, or not-for-profit sectors.

Data and code availability
The source code, frozen audited result artifacts, replay manifests, and analysis scripts are available at https://github.com/zhouz4884-coder/RevTrace.
""",
    )
    write_text(
        OUT / "09_editorial_inputs" / "suggested_reviewers_TEMPLATE.csv",
        "name,institution,email,expertise,conflict_check\n"
        "[REQUIRED],[REQUIRED],[REQUIRED],[REQUIRED],[confirm no conflict]\n",
    )
    write_text(
        OUT / "09_editorial_inputs" / "corresponding_phone_REQUIRED.txt",
        "Corresponding author: Xi Cao\nPhone: " +
        (phone if phone else "[REQUIRED BEFORE SUBMISSION]"),
    )
    write_text(
        OUT / "09_editorial_inputs" / "optional_graphical_abstract_README.md",
        """# Optional graphical abstract

The verified FGCS materials available to this package builder do not establish
that a graphical abstract is mandatory for every regular-paper submission.
If Editorial Manager requests one, use the corrected senior-author Figure 1
source and export a dedicated 1328x531-pixel-or-larger graphical abstract.
Do not upload an AI-generated graphical abstract.
""",
    )
    pending = [
        "Replace Figures 1--3 with the corrected senior-author exports listed in Paper/figures/SENIOR_FIGURE_CHANGE_REQUEST.md.",
        "Supply the corresponding-author phone number on the title page.",
        "Have all authors approve the final PDF, cover letter, and declarations.",
    ]
    write_text(
        OUT / "PENDING_BEFORE_UPLOAD.md",
        "# Pending before upload\n\n" + "\n".join(f"- [ ] {item}" for item in pending),
    )
    write_text(
        OUT / "SUBMISSION_CHECKLIST.md",
        """# FGCS submission checklist

- [x] Elsevier `elsarticle` double-column manuscript.
- [x] 18 pages including figures, tables, and references.
- [x] Abstract below 250 words; six English keywords.
- [x] Five highlights, each no more than 85 characters.
- [x] Cover letter and separate title page.
- [x] Competing-interest, AI-use, acknowledgments, and data statements.
- [x] Seven figures supplied as separate logically named files.
- [x] LaTeX source archive with editable tables and bibliography.
- [x] CJK replay, denominator, and contract-inventory audits pass.
- [ ] Corrected senior-author Figures 1--3 inserted and package rebuilt.
- [ ] Corresponding-author phone supplied.
- [ ] Suggested reviewers completed if requested by Editorial Manager.
- [ ] Final author approval and generated-PDF inspection.
""",
    )
    write_text(
        OUT / "README.md",
        """# RevTrace FGCS submission staging package

This directory is upload-structured but **not yet final**. It becomes ready
after the three corrected senior-author figures and the corresponding-author
phone number are supplied, the package builder is rerun, and the final PDF is
visually approved.

The manuscript is frozen at 18 pages. The package contains the manuscript PDF,
title page, cover letter, highlights, separate figures, LaTeX source zip,
declarations, audit evidence, and editorial-input templates.

Build command:

```powershell
python scripts/build_fgcs_submission_package.py --phone "<phone>"
```

The figure filenames referenced by LaTeX are stable, so corrected artwork can
replace the existing files without changing manuscript layout.
""",
    )
    write_text(
        OUT / "UPLOAD_MAPPING.csv",
        "editorial_manager_item,file,upload_status\n"
        "Manuscript,01_manuscript/RevTrace_FGCS_manuscript.pdf,ready after figure replacement\n"
        "Title Page,02_title_page/RevTrace_title_page.docx,pending phone\n"
        "Cover Letter,03_cover_letter/RevTrace_cover_letter.docx,ready\n"
        "Highlights,04_highlights/RevTrace_highlights.txt,ready\n"
        "Figure 1,05_figures/Figure_1.png,pending corrected export\n"
        "Figure 2,05_figures/Figure_2.jpg,pending corrected export\n"
        "Figure 3,05_figures/Figure_3.jpg,pending corrected export\n"
        "Figure 4,05_figures/Figure_4.pdf,ready\n"
        "Figure 5,05_figures/Figure_5.pdf,ready\n"
        "Figure 6,05_figures/Figure_6.pdf,ready\n"
        "Figure 7,05_figures/Figure_7.pdf,ready\n"
        "LaTeX Source,06_latex_source.zip,ready after figure replacement\n"
        "Declarations,08_declarations/declarations.txt,ready\n",
    )

    status = {
        "schema_version": 1,
        "decision": "STAGING_PENDING_FIGURES_AND_PHONE" if not phone else "STAGING_PENDING_FIGURES",
        "journal": "Future Generation Computer Systems",
        "article_type": "Regular research article",
        "manuscript_pages": 18,
        "highlights": len(highlights),
        "pending": pending if not phone else [pending[0], pending[2]],
    }
    write_text(OUT / "STATUS.json", json.dumps(status, indent=2))

    manifest_rows = []
    for path in sorted(OUT.rglob("*")):
        if path.is_file() and path.name not in {"MANIFEST.json", "SHA256SUMS.txt"}:
            manifest_rows.append({
                "path": path.relative_to(OUT).as_posix(),
                "bytes": path.stat().st_size,
                "sha256": sha256(path),
            })
    write_text(OUT / "MANIFEST.json", json.dumps({"files": manifest_rows}, indent=2))
    write_text(
        OUT / "SHA256SUMS.txt",
        "\n".join(f"{row['sha256']}  {row['path']}" for row in manifest_rows),
    )
    package_zip = OUT.parent / "FGCS_RevTrace_STAGING.zip"
    with zipfile.ZipFile(package_zip, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(OUT.rglob("*")):
            if path.is_file():
                archive.write(path, Path(OUT.name) / path.relative_to(OUT))
    print(json.dumps({
        "output": str(OUT),
        "package_zip": str(package_zip),
        "files": len(manifest_rows),
        "status": status["decision"],
    }))


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--phone", default=None)
    args = parser.parse_args()
    build_package(args.phone)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
