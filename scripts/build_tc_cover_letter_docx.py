"""Create the one-page IEEE TC cover letter as a submission-ready DOCX."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Paper_TC" / "submission" / "TC_COVER_LETTER.docx"


def set_font(run, *, size=11, bold=False, color="000000"):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_paragraph(doc, text="", *, bold=False, after=6, before=0, align=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.10
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_font(run, bold=bold)
    return paragraph


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run("COVER LETTER")
    set_font(run, size=16, bold=True, color="2E74B5")

    manuscript = doc.add_paragraph()
    manuscript.paragraph_format.space_after = Pt(14)
    run = manuscript.add_run(
        "RevTrace: A Transactional Runtime for Feedback-State Integrity in "
        "Revision-Generative ASR Serving"
    )
    set_font(run, size=12, bold=True)

    for line in (
        "3 September 2026",
        "Editor-in-Chief",
        "IEEE Transactions on Computers",
    ):
        add_paragraph(doc, line, after=0)
    add_paragraph(doc, "", after=6)
    add_paragraph(doc, "Dear Editor-in-Chief,", after=8)

    paragraphs = [
        "Please consider our manuscript, entitled \"RevTrace: A Transactional "
        "Runtime for Feedback-State Integrity in Revision-Generative ASR "
        "Serving,\" as a Regular Paper in IEEE Transactions on Computers.",
        "RevTrace is a runtime-systems contribution rather than a new ASR "
        "model. It identifies a feedback-state integrity failure caused when "
        "batch-dependent accelerator outputs are written into later execution "
        "state. To address this failure, the manuscript introduces "
        "Dual-Granularity Feedback Transactions (DGFT), a versioned state-write "
        "authorization protocol, and implements a parallel proposer-authority "
        "runtime. The evaluation covers transaction integrity, throughput, "
        "tail latency, energy, semantic quality, and replica scaling across "
        "Qwen3-ASR and Whisper on RTX 3090 and V100 GPU platforms.",
        "The manuscript makes three principal contributions. First, a causal "
        "study shows that a feedback edge can turn a local execution "
        "difference into persistent state divergence. Second, certificate, "
        "parent-version, frontier, and runtime-configuration binding provide "
        "formal accepted-history guarantees. Third, Coarse authority improves "
        "matched-workload throughput by 1.39x/1.28x; the complete two-plane "
        "system reaches up to 6.88x wall-clock speedup over its per-path "
        "all-singleton reference and scales to a balanced eight-V100 "
        "configuration.",
        "These contributions directly concern runtime systems, parallel "
        "accelerator execution, and the performance and reliability of "
        "computer systems. Streaming ASR is the latency-critical, "
        "revision-generative implementation and evaluation setting; the paper "
        "does not claim a new accelerator architecture, GPU kernel, "
        "datacenter-scale deployment, or multi-node system.",
        "This work is original, is not under consideration elsewhere, and has "
        "been approved by all authors. The authors declare no competing "
        "interests.",
    ]
    for text in paragraphs:
        add_paragraph(doc, text, after=7)

    add_paragraph(doc, "Sincerely,", after=8)
    add_paragraph(doc, "Xi Cao", bold=True, after=0)
    for line in (
        "Corresponding author",
        "Beijing University of Chemical Technology",
        "Beijing, China",
        "caoxi@buct.edu.cn",
    ):
        add_paragraph(doc, line, after=0)

    doc.core_properties.title = "Cover Letter - RevTrace - IEEE Transactions on Computers"
    doc.core_properties.author = "Xi Cao"
    doc.core_properties.subject = "Regular Paper submission"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
